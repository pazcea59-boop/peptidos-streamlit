"""
Registro de propiedades fisicoquímicas de péptidos (Streamlit)
================================================================
Aplicación para registrar, editar, guardar y exportar (Excel / PDF) un dataset
de péptidos.

Ejecución local:
    streamlit run app.py

Estructura de datos:
    - Tabla de trabajo (temporal): vive en st.session_state, se pierde al cerrar
      la pestaña o reiniciar la app.
    - Dataset guardado (persistente): data/peptidos_guardados.csv
"""
from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

st.set_page_config(page_title="Dataset de péptidos", page_icon="🧬", layout="wide")

# ---------------------------------------------------------------------------
# 1. CONSTANTES Y DATOS DE REFERENCIA
# ---------------------------------------------------------------------------
DATA_DIR = Path("data")
CSV_GUARDADO = DATA_DIR / "peptidos_guardados.csv"

COL_ID = "Identificador"
COLUMNS = [
    COL_ID,
    "Secuencia",
    "Longitud",
    "Masa molecular (Da)",
    "Polaridad",
    "% Apolares",
    "% Polares sin carga",
    "% Carga positiva",
    "% Carga negativa",
    "Composición de aminoácidos",
    "Hidrofobicidad (GRAVY)",
    "Carga neta aprox. (pH 7)",
    "Estructura 3D (PDB / enlace)",
    "Notas estructurales",
    "Propiedades complementarias",
    "Notas",
    "Fecha de registro",
]

# Escala de hidropatía de Kyte-Doolittle (1982)
KYTE_DOOLITTLE = {
    "A": 1.8, "R": -4.5, "N": -3.5, "D": -3.5, "C": 2.5,
    "Q": -3.5, "E": -3.5, "G": -0.4, "H": -3.2, "I": 4.5,
    "L": 3.8, "K": -3.9, "M": 1.9, "F": 2.8, "P": -1.6,
    "S": -0.8, "T": -0.7, "W": -0.9, "Y": -1.3, "V": 4.2,
}

# Masas promedio de residuos (Da); la masa del péptido suma además una molécula de agua
MASA_RESIDUO = {
    "A": 71.0788, "R": 156.1875, "N": 114.1038, "D": 115.0886, "C": 103.1388,
    "E": 129.1155, "Q": 128.1307, "G": 57.0519, "H": 137.1411, "I": 113.1594,
    "L": 113.1594, "K": 128.1741, "M": 131.1926, "F": 147.1766, "P": 97.1167,
    "S": 87.0782, "T": 101.1051, "W": 186.2132, "Y": 163.1760, "V": 99.1326,
}
MASA_AGUA = 18.01524

TRES_A_UNA = {
    "ala": "A", "arg": "R", "asn": "N", "asp": "D", "cys": "C",
    "gln": "Q", "glu": "E", "gly": "G", "his": "H", "ile": "I",
    "leu": "L", "lys": "K", "met": "M", "phe": "F", "pro": "P",
    "ser": "S", "thr": "T", "trp": "W", "tyr": "Y", "val": "V",
}

# Clasificación simplificada de residuos (heurística, ver README)
APOLARES = set("GAVLIMFWP")
POLARES_SIN_CARGA = set("STCYNQ")
POSITIVOS = set("KRH")
NEGATIVOS = set("DE")

FORMATO_1 = "Una letra (ej. ACDK)"
FORMATO_3 = "Tres letras (ej. Ala-Cys-Asp)"
AUTO = "Calcular automáticamente"
MANUAL = "Ingresar manualmente"
POLARIDADES = ["Polar", "Apolar (hidrofóbico)", "Anfipático / mixto",
               "Cargado positivo", "Cargado negativo"]

OPCIONES_MENU = [
    "Crear nuevo registro de péptido",
    "Gestionar tabla de trabajo (editar / guardar / exportar)",
    "Ver registros guardados (dataset)",
    "Importar registros desde CSV / Excel",
]

# Valores por defecto de los widgets del formulario (también sirven para limpiarlo)
DEFAULTS_FORM = {
    "f_id": "", "f_formato": FORMATO_1, "f_seq": "",
    "f_pol_modo": AUTO, "f_pol_manual": POLARIDADES[0],
    "f_comp_modo": AUTO, "f_comp_manual": "",
    "f_hid_modo": AUTO, "f_hid_manual": 0.0,
    "f_pdb": "", "f_estruc": "", "f_props": "", "f_notas": "",
}


# ---------------------------------------------------------------------------
# 2. LÓGICA DE PROPIEDADES FISICOQUÍMICAS
# ---------------------------------------------------------------------------
def parse_sequence(raw: str, formato: str) -> tuple[str | None, str | None]:
    """Normaliza una secuencia a código de una letra.

    Devuelve (secuencia, mensaje_de_error). Si hay error, secuencia es None.
    """
    limpio = re.sub(r"[\s\-_,.;]", "", raw or "")
    if not limpio:
        return None, "La secuencia está vacía."

    if formato == FORMATO_3:
        limpio = limpio.lower()
        if len(limpio) % 3 != 0:
            return None, "En código de tres letras, la longitud debe ser múltiplo de 3 (ej. AlaGly)."
        codigos = [limpio[i:i + 3] for i in range(0, len(limpio), 3)]
        invalidos = sorted({c for c in codigos if c not in TRES_A_UNA})
        if invalidos:
            return None, f"Códigos de tres letras no reconocidos: {', '.join(invalidos)}"
        return "".join(TRES_A_UNA[c] for c in codigos), None

    limpio = limpio.upper()
    invalidos = sorted(set(limpio) - set(KYTE_DOOLITTLE))
    if invalidos:
        return None, (f"Caracteres no válidos: {', '.join(invalidos)}. "
                      "Solo se admiten los 20 aminoácidos estándar.")
    return limpio, None


def compute_properties(seq: str) -> dict:
    """Calcula propiedades a partir de la secuencia (código de una letra)."""
    n = len(seq)
    cnt = Counter(seq)

    def pct(grupo: set) -> float:
        return round(100 * sum(cnt[a] for a in grupo) / n, 1)

    p_apolar = pct(APOLARES)
    if p_apolar >= 60:
        polaridad = "Apolar (hidrofóbico)"
    elif p_apolar <= 40:
        polaridad = "Polar"
    else:
        polaridad = "Anfipático / mixto"

    return {
        "longitud": n,
        "masa": round(sum(MASA_RESIDUO[a] * c for a, c in cnt.items()) + MASA_AGUA, 2),
        "gravy": round(sum(KYTE_DOOLITTLE[a] * c for a, c in cnt.items()) / n, 3),
        "p_apolar": p_apolar,
        "p_polar": pct(POLARES_SIN_CARGA),
        "p_pos": pct(POSITIVOS),
        "p_neg": pct(NEGATIVOS),
        "carga": cnt["K"] + cnt["R"] - cnt["D"] - cnt["E"],  # aproximación; ignora His y extremos
        "polaridad": polaridad,
        "composicion": "; ".join(f"{a}: {cnt[a]} ({100 * cnt[a] / n:.1f}%)" for a in sorted(cnt)),
    }


# ---------------------------------------------------------------------------
# 3. PERSISTENCIA (CSV) Y UTILIDADES DE TABLA
# ---------------------------------------------------------------------------
def init_state() -> None:
    if "registros" not in st.session_state:
        st.session_state.registros = pd.DataFrame(columns=COLUMNS)
    st.session_state.setdefault("editor_version", 0)
    for k, v in DEFAULTS_FORM.items():
        st.session_state.setdefault(k, v)


def flash(tipo: str, mensaje: str) -> None:
    """Guarda un mensaje para mostrarlo tras el siguiente rerun."""
    st.session_state["_flash"] = (tipo, mensaje)


def show_flash() -> None:
    if "_flash" in st.session_state:
        tipo, mensaje = st.session_state.pop("_flash")
        getattr(st, tipo)(mensaje)


def concat_seguro(frames: list[pd.DataFrame]) -> pd.DataFrame:
    frames = [f for f in frames if f is not None and not f.empty]
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=COLUMNS)


def limpiar(df: pd.DataFrame) -> pd.DataFrame:
    """Quita filas sin identificador y garantiza el orden de columnas."""
    df = df.copy()
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = None
    df = df[COLUMNS]
    mask = df[COL_ID].notna() & (df[COL_ID].astype(str).str.strip() != "")
    return df[mask].reset_index(drop=True)


def cargar_guardados() -> pd.DataFrame:
    if not CSV_GUARDADO.exists():
        return pd.DataFrame(columns=COLUMNS)
    df = pd.read_csv(CSV_GUARDADO, encoding="utf-8-sig")
    return limpiar(df)


def guardar_csv(df: pd.DataFrame) -> None:
    DATA_DIR.mkdir(exist_ok=True)
    df.to_csv(CSV_GUARDADO, index=False, encoding="utf-8-sig")


def fusionar_y_guardar(nuevos: pd.DataFrame, sobrescribir: bool) -> tuple[int, int, int]:
    """Añade `nuevos` al CSV. Devuelve (agregados, actualizados, omitidos)."""
    existentes = cargar_guardados()
    ya_existen = nuevos[COL_ID].astype(str).isin(set(existentes[COL_ID].astype(str)))
    n_dup = int(ya_existen.sum())
    if sobrescribir:
        combinado = concat_seguro([existentes, nuevos]).drop_duplicates(subset=COL_ID, keep="last")
        resultado = (len(nuevos) - n_dup, n_dup, 0)
    else:
        combinado = concat_seguro([existentes, nuevos[~ya_existen]])
        resultado = (len(nuevos) - n_dup, 0, n_dup)
    guardar_csv(combinado)
    return resultado


# ---------------------------------------------------------------------------
# 4. EXPORTACIÓN A EXCEL Y PDF (con caché para no regenerar en cada rerun)
# ---------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def build_excel(df: pd.DataFrame) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Peptidos")
        ws = writer.sheets["Peptidos"]

        relleno = PatternFill("solid", fgColor="1F4E78")
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for i, col in enumerate(df.columns, start=1):
            ancho = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str)])
            ws.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 60)

        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(vertical="top", wrap_text=True)

        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    return buffer.getvalue()


# Las fuentes estándar del PDF (Helvetica) solo cubren Latin-1; se sustituyen
# símbolos comunes en notas estructurales (α-hélice, β-lámina, etc.).
_SUSTITUCIONES_PDF = {
    "α": "alfa", "β": "beta", "γ": "gamma", "δ": "delta", "κ": "kappa",
    "−": "-", "–": "-", "—": "-", "≥": ">=", "≤": "<=", "→": "->",
    "µ": "u", "μ": "u",
}


def _pdf_txt(valor) -> str:
    """Convierte un valor en texto seguro para un Paragraph de ReportLab."""
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)) or str(valor).strip() == "":
        return "-"
    texto = str(valor)
    for origen, destino in _SUSTITUCIONES_PDF.items():
        texto = texto.replace(origen, destino)
    texto = texto.encode("cp1252", "replace").decode("cp1252")
    return escape(texto).replace("\n", "<br/>")


def _num(valor, formato: str) -> str:
    try:
        return format(float(valor), formato)
    except (TypeError, ValueError):
        return "-"


def _secuencia_en_bloques(seq, tam: int = 10) -> str:
    if seq is None or (not isinstance(seq, str)) or not seq:
        return "-"
    bloques = " ".join(seq[i:i + tam] for i in range(0, len(seq), tam))
    return f'<font name="Courier">{escape(bloques)}</font>'


@st.cache_data(show_spinner=False)
def build_pdf(df: pd.DataFrame) -> bytes:
    buffer = BytesIO()
    pagina = landscape(A4)
    doc = SimpleDocTemplate(
        buffer, pagesize=pagina, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.8 * cm,
        title="Dataset de péptidos", author="App de registro de péptidos",
    )
    estilos = getSampleStyleSheet()
    celda = ParagraphStyle("celda", parent=estilos["BodyText"], fontSize=8, leading=10)
    cabecera = ParagraphStyle("cabecera", parent=celda, fontName="Helvetica-Bold",
                              textColor=colors.white)

    historia = [
        Paragraph("Dataset de propiedades fisicoquímicas de péptidos", estilos["Title"]),
        Paragraph(f"Generado el {datetime.now():%Y-%m-%d %H:%M} &nbsp;|&nbsp; "
                  f"Total de registros: {len(df)}", estilos["Normal"]),
        Spacer(1, 0.5 * cm),
        Paragraph("Resumen", estilos["Heading2"]),
    ]

    # --- Tabla resumen -----------------------------------------------------
    encabezados = ["Identificador", "Longitud", "Masa (Da)", "Polaridad",
                   "GRAVY", "Carga neta", "PDB / enlace"]
    filas = [[Paragraph(h, cabecera) for h in encabezados]]
    for _, r in df.iterrows():
        filas.append([
            Paragraph(_pdf_txt(r[COL_ID]), celda),
            Paragraph(_num(r["Longitud"], ".0f"), celda),
            Paragraph(_num(r["Masa molecular (Da)"], ".2f"), celda),
            Paragraph(_pdf_txt(r["Polaridad"]), celda),
            Paragraph(_num(r["Hidrofobicidad (GRAVY)"], ".3f"), celda),
            Paragraph(_num(r["Carga neta aprox. (pH 7)"], ".0f"), celda),
            Paragraph(_pdf_txt(r["Estructura 3D (PDB / enlace)"]), celda),
        ])
    tabla = Table(filas, colWidths=[5 * cm, 3 * cm, 3 * cm, 4.5 * cm, 3 * cm, 3 * cm, 5.2 * cm],
                  repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B0B7BF")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    historia += [tabla, PageBreak(), Paragraph("Fichas detalladas", estilos["Heading2"])]

    # --- Una ficha por péptido --------------------------------------------
    campos = [c for c in COLUMNS if c != COL_ID]
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        datos = []
        for campo in campos:
            valor = (_secuencia_en_bloques(r[campo]) if campo == "Secuencia"
                     else _pdf_txt(r[campo]))
            datos.append([Paragraph(f"<b>{escape(campo)}</b>", celda), Paragraph(valor, celda)])
        ficha = Table(datos, colWidths=[6 * cm, 20.7 * cm])
        ficha.setStyle(TableStyle([
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C9CED4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        historia.append(KeepTogether([
            Paragraph(f"{i}. {_pdf_txt(r[COL_ID])}", estilos["Heading3"]),
            ficha,
            Spacer(1, 0.4 * cm),
        ]))

    def pie(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(pagina[0] - 1.5 * cm, 1 * cm, f"Página {documento.page}")
        canvas.restoreState()

    doc.build(historia, onFirstPage=pie, onLaterPages=pie)
    return buffer.getvalue()


def bloque_exportacion(df: pd.DataFrame, prefijo: str, nombre_base: str) -> None:
    """Botones de descarga de Excel, PDF y CSV."""
    st.subheader("Exportar")
    if df.empty:
        st.info("No hay registros para exportar.")
        return
    marca = datetime.now().strftime("%Y%m%d_%H%M")
    c1, c2, c3 = st.columns(3)
    c1.download_button(
        "📊 Exportar a Excel (.xlsx)", data=build_excel(df),
        file_name=f"{nombre_base}_{marca}.xlsx", key=f"{prefijo}_xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    c2.download_button(
        "📄 Exportar a PDF", data=build_pdf(df),
        file_name=f"{nombre_base}_{marca}.pdf", key=f"{prefijo}_pdf",
        mime="application/pdf",
    )
    c3.download_button(
        "🧾 Exportar a CSV", data=df.to_csv(index=False).encode("utf-8-sig"),
        file_name=f"{nombre_base}_{marca}.csv", key=f"{prefijo}_csv", mime="text/csv",
    )


# ---------------------------------------------------------------------------
# 5. PÁGINAS DE LA APLICACIÓN
# ---------------------------------------------------------------------------
def pantalla_bienvenida() -> None:
    st.markdown(
        "Registra péptidos con sus propiedades fisicoquímicas, acumúlalos en una tabla "
        "editable y expórtalos a **Excel** o **PDF**.\n\n"
        "Elige una opción en el menú de arriba para comenzar."
    )
    c1, c2 = st.columns(2)
    c1.metric("Registros en la tabla de trabajo", len(st.session_state.registros))
    c2.metric("Registros guardados en el dataset", len(cargar_guardados()))


def agregar_registro() -> None:
    """Callback del botón 'Agregar': valida, calcula y añade el registro."""
    s = st.session_state
    ident = s.f_id.strip()
    if not ident:
        flash("error", "El nombre / identificador es obligatorio.")
        return
    if ident in set(s.registros[COL_ID].astype(str)):
        flash("error", f"Ya existe un registro con el identificador «{ident}» en la tabla de trabajo.")
        return
    seq, error = parse_sequence(s.f_seq, s.f_formato)
    if error:
        flash("error", f"Secuencia inválida: {error}")
        return

    p = compute_properties(seq)
    comp_manual = s.f_comp_manual.strip()
    if s.f_comp_modo == MANUAL and not comp_manual:
        flash("error", "Has elegido composición manual, pero el campo está vacío.")
        return

    fila = {
        COL_ID: ident,
        "Secuencia": seq,
        "Longitud": p["longitud"],
        "Masa molecular (Da)": p["masa"],
        "Polaridad": p["polaridad"] if s.f_pol_modo == AUTO else s.f_pol_manual,
        "% Apolares": p["p_apolar"],
        "% Polares sin carga": p["p_polar"],
        "% Carga positiva": p["p_pos"],
        "% Carga negativa": p["p_neg"],
        "Composición de aminoácidos": p["composicion"] if s.f_comp_modo == AUTO else comp_manual,
        "Hidrofobicidad (GRAVY)": p["gravy"] if s.f_hid_modo == AUTO else float(s.f_hid_manual),
        "Carga neta aprox. (pH 7)": p["carga"],
        "Estructura 3D (PDB / enlace)": s.f_pdb.strip(),
        "Notas estructurales": s.f_estruc.strip(),
        "Propiedades complementarias": s.f_props.strip(),
        "Notas": s.f_notas.strip(),
        "Fecha de registro": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    s.registros = concat_seguro([s.registros, pd.DataFrame([fila])])
    flash("success", f"Registro «{ident}» agregado a la tabla de trabajo.")
    for k, v in DEFAULTS_FORM.items():  # limpiar formulario
        s[k] = v


def pagina_crear() -> None:
    st.header("Nuevo registro de péptido")
    show_flash()
    s = st.session_state

    # --- Identificación y secuencia ---------------------------------------
    st.subheader("1. Identificación y secuencia")
    st.text_input("Nombre / Identificador *", key="f_id",
                  placeholder="Ej. PEP-001 o Magainina 2")
    st.radio("Formato de la secuencia", [FORMATO_1, FORMATO_3], key="f_formato", horizontal=True)
    st.text_area("Secuencia de aminoácidos *", key="f_seq", height=100,
                 help="Se ignoran espacios, guiones y comas. Solo los 20 aminoácidos estándar.")

    seq, error = (None, None)
    if s.f_seq.strip():
        seq, error = parse_sequence(s.f_seq, s.f_formato)
    props = compute_properties(seq) if seq else None

    if error:
        st.warning(error)
    elif props:
        st.caption("Vista previa de valores calculados a partir de la secuencia")
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Longitud", props["longitud"])
        m2.metric("Masa (Da)", f"{props['masa']:.2f}")
        m3.metric("GRAVY", f"{props['gravy']:.3f}")
        m4.metric("Carga neta aprox.", f"{props['carga']:+d}")
        st.caption(f"Secuencia normalizada (una letra): {seq}")

    # --- Propiedades fisicoquímicas ---------------------------------------
    st.subheader("2. Propiedades fisicoquímicas")
    col_pol, col_comp, col_hid = st.columns(3)

    with col_pol:
        st.radio("Polaridad", [AUTO, MANUAL], key="f_pol_modo")
        if s.f_pol_modo == AUTO:
            st.info(props["polaridad"] if props else "Se calculará con una secuencia válida.")
        else:
            st.selectbox("Polaridad (manual)", POLARIDADES, key="f_pol_manual")

    with col_comp:
        st.radio("Composición de aminoácidos", [AUTO, MANUAL], key="f_comp_modo")
        if s.f_comp_modo == AUTO:
            st.info(props["composicion"] if props else "Se calculará con una secuencia válida.")
        else:
            st.text_area("Composición (manual)", key="f_comp_manual", height=100,
                         placeholder="Ej. A: 3 (15%); L: 5 (25%) ...")

    with col_hid:
        st.radio("Hidrofobicidad (índice hidropático)", [AUTO, MANUAL], key="f_hid_modo")
        if s.f_hid_modo == AUTO:
            st.info(f"GRAVY (Kyte-Doolittle): {props['gravy']:.3f}" if props
                    else "Se calculará con una secuencia válida.")
        else:
            st.number_input("Índice hidropático (manual)", key="f_hid_manual", step=0.01, format="%.3f")

    # --- Estructura 3D -----------------------------------------------------
    st.subheader("3. Estructura tridimensional")
    st.text_input("ID o enlace a archivo PDB (opcional)", key="f_pdb",
                  placeholder="Ej. 2MAG o https://www.rcsb.org/structure/2MAG")
    st.text_area("Descripción / notas estructurales (opcional)", key="f_estruc", height=80,
                 placeholder="Ej. α-hélice anfipática en medio lipídico; método RMN...")

    # --- Campos adicionales -----------------------------------------------
    st.subheader("4. Información complementaria")
    st.text_area("Propiedades complementarias (opcional)", key="f_props", height=80,
                 placeholder="Ej. punto isoeléctrico, solubilidad, actividad biológica...")
    st.text_area("Notas (opcional)", key="f_notas", height=80)

    st.button("➕ Agregar registro a la tabla de trabajo", type="primary", on_click=agregar_registro)

    if not s.registros.empty:
        st.divider()
        st.caption(f"Últimos registros en la tabla de trabajo ({len(s.registros)} en total)")
        st.dataframe(s.registros.tail(5), width="stretch", hide_index=True)


def pagina_gestionar() -> None:
    st.header("Tabla de trabajo")
    show_flash()
    s = st.session_state
    if s.registros.empty:
        st.info("La tabla de trabajo está vacía. Crea registros o importa un archivo.")
        return

    st.caption("Doble clic en una celda para editar. Selecciona filas y pulsa Supr para eliminarlas. "
               "Recuerda pulsar «Aplicar cambios» para conservar lo editado.")
    editado = st.data_editor(
        s.registros, num_rows="dynamic", width="stretch", hide_index=True,
        key=f"editor_{s.editor_version}",
    )
    editado = limpiar(editado)
    if editado[COL_ID].duplicated().any():
        st.warning("Hay identificadores duplicados en la tabla.")

    c1, c2 = st.columns(2)
    if c1.button("✅ Aplicar cambios a la tabla"):
        s.registros = editado
        s.editor_version += 1  # reinicia el editor para evitar ediciones "fantasma"
        flash("success", "Cambios aplicados.")
        st.rerun()

    with c2:
        sobrescribir = st.checkbox("Sobrescribir registros con el mismo identificador al guardar")
        if st.button("💾 Guardar definitivamente en el dataset (CSV)"):
            s.registros = editado
            nuevos, actualizados, omitidos = fusionar_y_guardar(editado, sobrescribir)
            s.editor_version += 1
            flash("success", f"Dataset guardado: {nuevos} nuevos, {actualizados} actualizados, "
                             f"{omitidos} omitidos por ID repetido.")
            st.rerun()

    with st.expander("Vaciar tabla de trabajo"):
        if st.checkbox("Confirmo que quiero eliminar todos los registros de la tabla de trabajo"):
            if st.button("🗑️ Vaciar tabla"):
                s.registros = pd.DataFrame(columns=COLUMNS)
                s.editor_version += 1
                st.rerun()

    st.divider()
    bloque_exportacion(editado, "trabajo", "peptidos_tabla_trabajo")


def pagina_guardados() -> None:
    st.header("Registros guardados en el dataset")
    show_flash()
    guardados = cargar_guardados()
    if guardados.empty:
        st.info("Aún no hay registros guardados. Guarda desde la tabla de trabajo.")
        return

    filtro = st.text_input("Buscar", placeholder="Identificador, secuencia, notas...")
    vista = guardados
    if filtro:
        mascara = guardados.astype(str).apply(
            lambda col: col.str.contains(filtro, case=False, na=False, regex=False)
        ).any(axis=1)
        vista = guardados[mascara]
    st.caption(f"Mostrando {len(vista)} de {len(guardados)} registros")
    st.dataframe(vista, width="stretch", hide_index=True)

    c1, c2 = st.columns(2)
    if c1.button("⬆️ Cargar registros mostrados en la tabla de trabajo"):
        actuales = set(st.session_state.registros[COL_ID].astype(str))
        nuevos = vista[~vista[COL_ID].astype(str).isin(actuales)]
        st.session_state.registros = concat_seguro([st.session_state.registros, nuevos])
        st.session_state.editor_version += 1
        flash("success", f"{len(nuevos)} registros cargados en la tabla de trabajo.")
        st.rerun()

    with c2.expander("Eliminar registros del dataset guardado"):
        ids = st.multiselect("Identificadores a eliminar", guardados[COL_ID].astype(str).tolist())
        confirmar = st.checkbox("Confirmo la eliminación permanente")
        if st.button("🗑️ Eliminar seleccionados") and ids and confirmar:
            guardar_csv(guardados[~guardados[COL_ID].astype(str).isin(ids)])
            flash("success", f"{len(ids)} registros eliminados.")
            st.rerun()

    st.divider()
    st.caption("Se exporta lo que se muestra en la tabla (aplica el filtro de búsqueda).")
    bloque_exportacion(vista, "guardados", "peptidos_dataset")


def pagina_importar() -> None:
    st.header("Importar registros")
    st.caption(f"El archivo debe contener al menos la columna «{COL_ID}». "
               "Las columnas ausentes se dejarán vacías. Lo ideal es importar un archivo "
               "exportado previamente por esta app.")
    archivo = st.file_uploader("Archivo CSV o Excel", type=["csv", "xlsx"])
    if archivo is None:
        return
    try:
        df = (pd.read_csv(archivo, encoding="utf-8-sig") if archivo.name.lower().endswith(".csv")
              else pd.read_excel(archivo))
    except Exception as exc:  # noqa: BLE001 - mostrar cualquier error de lectura al usuario
        st.error(f"No se pudo leer el archivo: {exc}")
        return
    if COL_ID not in df.columns:
        st.error(f"No se encontró la columna «{COL_ID}».")
        return

    df = limpiar(df)
    st.dataframe(df, width="stretch", hide_index=True)
    if st.button("Agregar a la tabla de trabajo", type="primary"):
        actuales = set(st.session_state.registros[COL_ID].astype(str))
        nuevos = df[~df[COL_ID].astype(str).isin(actuales)]
        st.session_state.registros = concat_seguro([st.session_state.registros, nuevos])
        st.session_state.editor_version += 1
        st.success(f"{len(nuevos)} registros agregados; "
                   f"{len(df) - len(nuevos)} omitidos por identificador repetido.")


# ---------------------------------------------------------------------------
# 6. PUNTO DE ENTRADA
# ---------------------------------------------------------------------------
def main() -> None:
    init_state()
    st.title("🧬 Dataset de propiedades fisicoquímicas de péptidos")

    opcion = st.selectbox("¿Qué deseas realizar?", OPCIONES_MENU, index=None,
                          placeholder="Selecciona una opción…", key="menu")

    with st.sidebar:
        st.subheader("Estado")
        st.metric("Tabla de trabajo", len(st.session_state.registros))
        st.caption("La tabla de trabajo es temporal: guárdala en el dataset para conservarla.")

    st.divider()
    if opcion is None:
        pantalla_bienvenida()
    elif opcion == OPCIONES_MENU[0]:
        pagina_crear()
    elif opcion == OPCIONES_MENU[1]:
        pagina_gestionar()
    elif opcion == OPCIONES_MENU[2]:
        pagina_guardados()
    else:
        pagina_importar()


main()
